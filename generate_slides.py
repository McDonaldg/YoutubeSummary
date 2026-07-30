"""構造化Markdownからプレゼンスライド（pptx / HTML）を生成する。"""
from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

import config


@dataclass
class Slide:
    heading: str
    bullets: list[str]


@dataclass
class Deck:
    title: str
    meta: dict[str, str]
    slides: list[Slide]


_FRONTMATTER_RE = re.compile(r"^---\n(.*?)\n---\n\n?", re.DOTALL)


def _parse_frontmatter(text: str) -> tuple[dict[str, str], str]:
    match = _FRONTMATTER_RE.match(text)
    if not match:
        return {}, text
    meta: dict[str, str] = {}
    for line in match.group(1).splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            meta[k.strip()] = v.strip().strip('"')
    return meta, text[match.end() :]


def parse_markdown(text: str) -> Deck:
    meta, body = _parse_frontmatter(text)

    title_match = re.match(r"^#\s+(.+)$", body, re.MULTILINE)
    title = title_match.group(1).strip() if title_match else meta.get("title", "Untitled")
    body = body[title_match.end() :] if title_match else body

    slides: list[Slide] = []
    sections = re.split(r"^##\s+", body, flags=re.MULTILINE)[1:]
    for section in sections:
        lines = section.strip("\n").splitlines()
        heading = lines[0].strip()
        bullets = [
            re.sub(r"^[-*]\s+", "", line).strip()
            for line in lines[1:]
            if line.strip().startswith(("-", "*"))
        ]
        if heading:
            slides.append(Slide(heading=heading, bullets=bullets))

    return Deck(title=title, meta=meta, slides=slides)


def generate_pptx(deck: Deck, out_path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()

    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = deck.title
    subtitle_parts = [p for p in (deck.meta.get("channel"), deck.meta.get("url")) if p]
    try:
        slide.placeholders[1].text = " / ".join(subtitle_parts)
    except KeyError:
        pass

    bullet_layout = prs.slide_layouts[1]
    for section in deck.slides:
        s = prs.slides.add_slide(bullet_layout)
        s.shapes.title.text = section.heading
        body = s.placeholders[1].text_frame
        body.clear()
        if not section.bullets:
            body.text = "(内容なし)"
        else:
            body.text = section.bullets[0]
            for bullet in section.bullets[1:]:
                p = body.add_paragraph()
                p.text = bullet
                p.level = 0
            for p in body.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(20)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="ja">
<head>
<meta charset="UTF-8">
<title>{title}</title>
<style>
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; font-family: "Hiragino Sans", "Noto Sans JP", sans-serif; background: #1a1a2e; }}
  .slide {{
    display: none; width: 100vw; height: 100vh; padding: 8vh 8vw;
    color: #eee; background: linear-gradient(135deg, #1a1a2e, #16213e);
  }}
  .slide.active {{ display: flex; flex-direction: column; justify-content: center; }}
  .slide h1 {{ font-size: 3rem; margin-bottom: 1rem; }}
  .slide h2 {{ font-size: 2.5rem; color: #4dd0e1; margin-bottom: 1.5rem; }}
  .slide ul {{ font-size: 1.6rem; line-height: 2.2; }}
  .meta {{ color: #aaa; font-size: 1.1rem; margin-top: 2rem; }}
  .nav {{
    position: fixed; bottom: 1rem; right: 1.5rem; color: #888; font-size: 0.9rem;
  }}
</style>
</head>
<body>
{slides}
<div class="nav">&larr; / &rarr; キーでスライド移動 (<span id="cur">1</span>/<span id="total"></span>)</div>
<script>
  const slides = document.querySelectorAll('.slide');
  let i = 0;
  document.getElementById('total').textContent = slides.length;
  function show(n) {{
    slides[i].classList.remove('active');
    i = (n + slides.length) % slides.length;
    slides[i].classList.add('active');
    document.getElementById('cur').textContent = i + 1;
  }}
  document.addEventListener('keydown', (e) => {{
    if (e.key === 'ArrowRight' || e.key === ' ') show(i + 1);
    if (e.key === 'ArrowLeft') show(i - 1);
  }});
  show(0);
</script>
</body>
</html>
"""


def generate_html(deck: Deck, out_path: Path) -> Path:
    slides_html = []

    subtitle_parts = [p for p in (deck.meta.get("channel"), deck.meta.get("url")) if p]
    slides_html.append(
        f'<section class="slide active"><h1>{deck.title}</h1>'
        f'<p class="meta">{" / ".join(subtitle_parts)}</p></section>'
    )

    for section in deck.slides:
        items = "".join(f"<li>{b}</li>" for b in section.bullets) or "<li>(内容なし)</li>"
        slides_html.append(f'<section class="slide"><h2>{section.heading}</h2><ul>{items}</ul></section>')

    html = _HTML_TEMPLATE.format(title=deck.title, slides="\n".join(slides_html))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html, encoding="utf-8")
    return out_path


def generate(video_id: str, fmt: str | None = None) -> Path:
    fmt = (fmt or config.SLIDE_FORMAT).lower()
    if fmt not in ("pptx", "html"):
        raise ValueError(f"未対応のフォーマットです: {fmt}（pptx または html）")

    md_path = config.OUTPUT_MARKDOWN_DIR / f"{video_id}.md"
    if not md_path.exists():
        raise FileNotFoundError(f"要約Markdownが見つかりません: {md_path}")

    deck = parse_markdown(md_path.read_text(encoding="utf-8"))
    out_path = config.OUTPUT_SLIDES_DIR / f"{video_id}.{fmt}"

    if fmt == "pptx":
        generate_pptx(deck, out_path)
    else:
        generate_html(deck, out_path)

    print(f"[generate_slides] 生成完了: {out_path} ({len(deck.slides)}スライド + タイトル)")
    return out_path


def _main() -> None:
    parser = argparse.ArgumentParser(description="Markdownからスライドを生成")
    parser.add_argument("--video-id", required=True)
    parser.add_argument("--format", choices=["pptx", "html"], default=None)
    args = parser.parse_args()
    generate(args.video_id, fmt=args.format)


if __name__ == "__main__":
    _main()
